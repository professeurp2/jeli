"""Documents: what members shared in the groups, or the team added on the dashboard — PDF, Word, text.

Jeli keeps the file itself, and learns its text page by page (a document's pages are messages of
chat_id = the document's id), so answers quote "page 3" and /search finds it. When a member asks for
a document, Jeli sends the file, as WhatsApp does; asked for it in another language, it translates
it into a PDF — sent a minute later — and keeps that translation for the next request.
"""

import asyncio
import hashlib
import io
import logging
import re
import unicodedata
from datetime import datetime, timedelta, timezone

from pydantic import BaseModel

from app.answer.citations import display_author, short_day
from app.answer.conversation import Turn
from app.answer.language import TEXTS, detect_language
from app.answer.llm import LLM, LLMUnavailable
from app.answer.understand import conversation_text
from app.kb.indexer import DOCUMENT_PREFIX, document_page
from app.kb.store import Store
from app.models import Attachment, Document, Reply, StoredMessage

log = logging.getLogger(__name__)

MAX_BYTES = 15 * 1024 * 1024
PAGE_CHARS = 2500  # Word and text files have no pages: parts of about this size stand for them
PART_CHARS = 1200  # a page is stored in parts of at most this size, the unit of retrieval
TRANSLATION_BATCH_CHARS = 5000
TRANSLATION_TIMEOUT = 120
TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}
# Languages the PDF fonts can write (Latin script).
LANGUAGE_NAMES = {
    "en": "English", "fr": "French", "pt": "Portuguese", "es": "Spanish", "sw": "Swahili",
    "de": "German", "it": "Italian", "nl": "Dutch",
}
NAMES_FR = {
    "en": "anglais", "fr": "français", "pt": "portugais", "es": "espagnol", "sw": "swahili",
    "de": "allemand", "it": "italien", "nl": "néerlandais",
}
# Each language in its own words, for the translated document's title.
NATIVE_NAMES = {
    "en": "English", "fr": "français", "pt": "português", "es": "español", "sw": "Kiswahili",
    "de": "Deutsch", "it": "italiano", "nl": "Nederlands",
}


def extension(filename: str, mimetype: str = "") -> str | None:
    """The kind of document, from its name or its type; None when Jeli cannot read it."""
    name = filename.lower()
    for ext, mime in TYPES.items():
        if name.endswith(ext) or (mimetype and mimetype.split(";")[0].strip() == mime):
            return ext
    return None


def _split(text: str, size: int) -> list[str]:
    """Parts of at most `size` characters, cut between paragraphs, else between sentences."""
    parts, current = [], ""
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        while len(paragraph) > size:
            cut = paragraph.rfind(". ", 0, size)
            cut = cut + 1 if cut > size // 3 else size
            if current:
                parts.append(current)
                current = ""
            parts.append(paragraph[:cut].strip())
            paragraph = paragraph[cut:].strip()
        if current and len(current) + len(paragraph) + 2 > size:
            parts.append(current)
            current = ""
        current = f"{current}\n\n{paragraph}" if current else paragraph
    if current:
        parts.append(current)
    return parts


def read_pages(ext: str, data: bytes) -> list[str]:
    """The text of each page (PDF), or of each part of about a page (Word, text)."""
    if ext == ".pdf":
        from pypdf import PdfReader

        return [(page.extract_text() or "").strip() for page in PdfReader(io.BytesIO(data)).pages]
    if ext == ".docx":
        from docx import Document as Docx

        text = "\n\n".join(p.text for p in Docx(io.BytesIO(data)).paragraphs if p.text.strip())
        return _split(text, PAGE_CHARS)
    return _split(data.decode("utf-8-sig", errors="replace"), PAGE_CHARS)


def title_of(filename: str) -> str:
    stem = filename.rsplit(".", 1)[0]
    return " ".join(re.sub(r"[_]+", " ", stem).split())[:120] or "Document"


def slugify(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_text.lower()).strip("-")[:40] or "document"


def page_messages(document: Document, pages: list[str]) -> list[StoredMessage]:
    """A document's text as messages: page p at shared_at + p seconds, in parts."""
    messages = []
    for number, text in enumerate(pages, 1):
        for part, chunk in enumerate(_split(text, PART_CHARS)):
            messages.append(
                StoredMessage(
                    id=f"{document.id}#p{number}.{part}",
                    chat_id=document.id,
                    source="document",
                    author=document.shared_by or "Team",
                    sent_at=document.shared_at + timedelta(seconds=number, microseconds=part),
                    text=chunk,
                )
            )
    return messages


FILE_NAME = re.compile(r"([\w\-. ()'’&,]+?\.(?:pdf|docx?|pptx?|xlsx?|txt))", re.IGNORECASE)
MARKERS = re.compile(r"<document omis>|document omitted|<attached:[^>]*>|\(fichier joint\)|\(file attached\)|<pièce jointe[^>]*>|\[transféré\]|\[forwarded\]", re.IGNORECASE)


async def missing_documents(store: Store) -> list[dict]:
    """Documents members shared in the groups whose file Jeli does not have (chat histories
    exported without media): name, who shared it, when, where."""
    kept = await store.list_documents()
    known = {d.filename.lower() for d in kept} | {d.title.lower() for d in kept}
    missing, seen = [], set()
    for row in await store.mentioned_documents():
        found = FILE_NAME.search(row["text"])
        name = found.group(1).strip() if found else " ".join(MARKERS.sub("", row["text"]).split())[:80]
        plain = re.sub(r"^\d{6,}-", "", name)
        if not name or plain.lower() in known or plain.rsplit(".", 1)[0].lower() in known or plain.lower() in seen:
            continue
        seen.add(plain.lower())
        missing.append({"name": plain, "author": row["author"], "sent_at": row["sent_at"], "chat_id": row["chat_id"]})
    return missing


class FileChoice(BaseModel):
    document: int
    translate_to: str


class Block(BaseModel):
    kind: str
    text: str


class Translated(BaseModel):
    blocks: list[Block]
    title: str = ""
    note: str = ""


CHOOSE_SYSTEM = """\
A member of a WhatsApp community asks Jeli, its assistant, for a file. From their message and the
conversation, decide which of the listed documents they want: "document" is its number, or 0 when
they ask for none of these (for example a question about a topic, or a document not in the list).
"translate_to" is the two-letter code of the language they want it in (en, fr, pt, es, sw, …) when
they ask for a translation or a version in another language; otherwise "".
"""

# Said plainly, in the system and in the request: with a softer wording the light models return the
# text unchanged (measured: 95% of the words copied, in French as in Swahili).
TRANSLATE_SYSTEM = """\
You are a professional translator. Translate the document the user sends from its language into
{language}: every heading, sentence and list item must be written in {language}. Keep names, numbers,
dates, amounts and links exactly, and technical words readers use as they are (hackathon, chatbot,
AI, demo, pitch); never add, drop or summarise anything. Return it as blocks, in order: "heading"
for titles, "bullet" for items of a list, "paragraph" for the rest.
"""
TRANSLATE_REQUEST = "Translate into {language}:\n\n{text}"
TRANSLATE_FIRST = (
    "\n\nAlso return \"title\": the title «{title}» in {language}, and \"note\": this sentence in {language}: «{note}»"
)
# A part whose words are mostly those of the original was not translated.
COPIED_SHARE = 0.5
COPIED_MIN_WORDS = 30


def copied(source: str, translated: str) -> bool:
    """True when `translated` is, in fact, `source` again (the model did not translate)."""
    vocabulary = set(re.findall(r"[^\W\d_]{4,}", source.lower()))
    words = re.findall(r"[^\W\d_]{4,}", translated.lower())
    if len(words) < COPIED_MIN_WORDS:
        return False  # too short to tell: names, figures, a title
    return sum(word in vocabulary for word in words) / len(words) > COPIED_SHARE


class NotTranslated(LLMUnavailable):
    """The model returned the original instead of a translation."""


class Documents:
    def __init__(self, store: Store, llm: LLM | None):
        self.store = store
        self.llm = llm
        self.known_names: dict[str, str] = {}  # number → name, given by the team

    def who(self, author: str) -> str:
        return self.known_names.get(re.sub(r"\D", "", author)) or display_author(author) if author else "a member"

    async def add(
        self,
        filename: str,
        data: bytes,
        *,
        title: str = "",
        shared_by: str = "Team",
        shared_at: datetime | None = None,
        chat_id: str = "",
        mimetype: str = "",
    ) -> tuple[Document, bool]:
        """Keep a document and learn its text. Returns it, and whether it was new.
        ValueError, in words for the team, when it cannot be read."""
        ext = extension(filename, mimetype)
        if ext is None:
            raise ValueError("Jeli reads PDF, Word (.docx) and text files")
        if len(data) > MAX_BYTES:
            raise ValueError("this file is too big (15 MB at most)")
        try:
            pages = await asyncio.to_thread(read_pages, ext, data)
        except Exception as error:  # a damaged or protected file
            raise ValueError("this file cannot be opened") from error
        if not any(page.strip() for page in pages):
            raise ValueError("no text found in this file (a scan?)")
        title = " ".join(title.split())[:120] or title_of(filename)
        shared_at = shared_at or datetime.now(timezone.utc)
        document = Document(
            id=f"{DOCUMENT_PREFIX}{slugify(title)}-{hashlib.sha256(data).hexdigest()[:10]}",
            title=title,
            filename=filename[:160],
            mimetype=TYPES[ext],
            size_bytes=len(data),
            pages=len(pages),
            language=detect_language(" ".join(pages)[:3000]),
            shared_by=shared_by,
            shared_at=shared_at,
            chat_id=chat_id,
        )
        new = await self.store.save_document(document, data)
        if new:
            await self.store.add_messages(page_messages(document, pages))
            log.info("Document %s kept: %d pages", document.title, document.pages)
        return document, new

    async def attachment(self, document: Document, caption: str = "") -> Attachment | None:
        data = await self.store.document_content(document.id)
        if data is None:
            return None
        return Attachment(document.filename, document.mimetype, data, caption=caption, document_id=document.id)

    async def reply(self, text: str, language: str, turns: list[Turn] = ()) -> Reply | None:
        """A member asks for a document, maybe translated: the file (at once, or once translated),
        or None when they ask for none of the documents Jeli keeps."""
        documents = await self.store.list_documents()
        if not documents or self.llm is None:
            return None
        listing = "\n".join(
            f"[{n}] «{d.title}» ({d.filename}, {d.pages} pages, in {LANGUAGE_NAMES.get(d.language, d.language)}), "
            f"shared by {self.who(d.shared_by)} on {short_day(d.shared_at)}"
            for n, d in enumerate(documents[:40], 1)
        )
        prompt = (
            f"Documents:\n{listing}\n\n"
            + (f"Conversation so far:\n{conversation_text(list(turns))}\n\n" if turns else "")
            + f"Latest message:\n{text}"
        )
        try:
            choice = await self.llm.generate(prompt, FileChoice, system=CHOOSE_SYSTEM, temperature=0, attempts=2)
        except LLMUnavailable:
            return None
        if not 1 <= choice.document <= min(len(documents), 40):
            return None
        document = documents[choice.document - 1]
        texts = TEXTS[language]
        target = choice.translate_to.strip().lower()[:2]
        who, day = self.who(document.shared_by), short_day(document.shared_at)
        if not target or target == document.language:
            file = await self.attachment(document, caption=document.title)
            if file:
                return Reply(texts["file_here"].format(title=document.title, who=who, day=day), attachment=file)
            # Document is indexed (Jeli can answer questions about it) but the file itself was not
            # stored — typically a chat export that included text but not the attached file.
            return Reply(
                f"📄 I've read «{document.title}» and can answer questions about it, "
                f"but I don't have the file to send. Ask {who} to share it again in the group!"
                if language == "en" else
                f"📄 J'ai lu «{document.title}» et je peux répondre à vos questions dessus, "
                f"mais je n'ai pas le fichier pour l'envoyer. Demandez à {who} de le repartager dans le groupe !"
            )
        if target not in LANGUAGE_NAMES:
            names = LANGUAGE_NAMES if language == "en" else NAMES_FR
            return Reply(texts["file_language_unsupported"].format(languages=", ".join(names.values())))
        language_name = (LANGUAGE_NAMES if language == "en" else NAMES_FR)[target]
        cached = await self.store.translation(document.id, target)
        if cached:
            file = await self.attachment(cached, caption=cached.title)
            return Reply(texts["file_here"].format(title=cached.title, who="Jeli", day=short_day(cached.shared_at)), attachment=file)

        async def translation() -> Attachment | str:
            try:
                return await self.translate(document, target)
            except LLMUnavailable:
                return texts["file_translate_failed"].format(title=document.title)

        return Reply(texts["file_translating"].format(title=document.title, language=language_name), pending=translation)

    async def translate(self, document: Document, target: str) -> Attachment:
        """The document in another language, as a PDF, kept for the next request."""
        from app.answer.pdf import build_pdf

        messages = await self.store.messages_of(document.id)
        pages: dict[int, list[str]] = {}
        for message in messages:
            pages.setdefault(document_page(message.sent_at, document.shared_at), []).append(message.text)
        texts = ["\n\n".join(parts) for _, parts in sorted(pages.items())]
        batches, current = [], ""
        for page in texts:
            if current and len(current) + len(page) > TRANSLATION_BATCH_CHARS:
                batches.append(current)
                current = ""
            current = f"{current}\n\n{page}" if current else page
        if current:
            batches.append(current)
        name, native = LANGUAGE_NAMES[target], NATIVE_NAMES[target]
        title = f"{document.title} ({native})"
        note = (
            f"Machine translation by Jeli of «{document.title}», shared by {self.who(document.shared_by)} "
            f"on {short_day(document.shared_at)}. The original prevails."
        )
        heading, foreword = "", ""
        blocks: list[tuple[str, str]] = []
        for number, batch in enumerate(batches):
            request = TRANSLATE_REQUEST.format(language=name, text=batch)
            if number == 0:
                request += TRANSLATE_FIRST.format(title=document.title, language=name, note=note)
            translated = await self._translate_part(batch, request, name)
            heading, foreword = heading or translated.title, foreword or translated.note
            blocks += [(b.kind if b.kind in ("heading", "bullet") else "paragraph", b.text) for b in translated.blocks]
        if heading.strip():
            title = f"{heading.strip()} ({native})"
        pdf = await asyncio.to_thread(build_pdf, title, foreword.strip() or note, blocks)
        stem = document.filename.rsplit(".", 1)[0]
        translation = Document(
            id=f"{document.id}~{target}",
            title=title,
            filename=f"{stem} ({name}).pdf",
            mimetype=TYPES[".pdf"],
            size_bytes=len(pdf),
            pages=0,
            language=target,
            shared_by="Jeli",
            shared_at=datetime.now(timezone.utc),
            chat_id=document.chat_id,
            translation_of=document.id,
        )
        await self.store.save_document(translation, pdf)
        log.info("Translated %s into %s", document.title, name)
        return Attachment(translation.filename, translation.mimetype, pdf, caption=title, document_id=translation.id)

    async def _translate_part(self, source: str, request: str, language: str) -> Translated:
        """One part, translated — asked twice at most: a copy of the original is never kept."""
        for _ in range(2):
            translated = await self.llm.generate(
                request, Translated, system=TRANSLATE_SYSTEM.format(language=language), timeout=TRANSLATION_TIMEOUT, temperature=0
            )
            if not copied(source, " ".join(b.text for b in translated.blocks)):
                return translated
            log.warning("The model returned the original instead of %s", language)
        raise NotTranslated
